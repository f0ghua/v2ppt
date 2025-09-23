"""
PowerPoint generation module for creating PPT files from extracted images
"""
import logging
from pathlib import Path
from typing import List, Optional, Tuple
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from PIL import Image
import time

from config.settings import PPTSettings
from src.utils import ProgressTracker, format_file_size

logger = logging.getLogger('v2ppt.ppt')

class PPTGenerator:
    """
    PowerPoint generator for creating presentations from images
    """
    
    def __init__(self, template_path: Optional[Path] = None):
        """
        Initialize PPT generator
        
        Args:
            template_path: Path to template PPTX file (optional)
        """
        self.template_path = template_path
        self.presentation = None
        
        # Initialize presentation
        self._initialize_presentation()
    
    def _initialize_presentation(self):
        """Initialize PowerPoint presentation"""
        if self.template_path and self.template_path.exists():
            # Load from template
            self.presentation = Presentation(str(self.template_path))
            logger.info(f"Loaded template: {self.template_path.name}")
        else:
            # Create new presentation
            self.presentation = Presentation()
            
            # Set slide dimensions
            self.presentation.slide_width = Inches(PPTSettings.SLIDE_WIDTH)
            self.presentation.slide_height = Inches(PPTSettings.SLIDE_HEIGHT)
            
            logger.info("Created new presentation")
            logger.info(f"Slide dimensions: {PPTSettings.SLIDE_WIDTH}\" x {PPTSettings.SLIDE_HEIGHT}\"")
    
    def add_image_slide(self, image_path: Path, slide_title: Optional[str] = None) -> bool:
        """
        Add a slide with an image
        
        Args:
            image_path: Path to image file
            slide_title: Optional slide title
            
        Returns:
            True if successful
        """
        try:
            # Validate image
            if not image_path.exists():
                logger.error(f"Image file not found: {image_path}")
                return False
            
            # Get image dimensions for optimal positioning
            img_width, img_height = self._get_image_dimensions(image_path)
            
            # Add slide with blank layout
            slide_layout = self.presentation.slide_layouts[6]  # Blank layout
            slide = self.presentation.slides.add_slide(slide_layout)
            
            # Add title if provided
            if slide_title:
                self._add_slide_title(slide, slide_title)
            
            # Calculate image position and size
            left, top, width, height = self._calculate_image_position(img_width, img_height)
            
            # Add image to slide
            slide.shapes.add_picture(
                str(image_path),
                left=Inches(left),
                top=Inches(top),
                width=Inches(width),
                height=Inches(height)
            )
            
            logger.debug(f"Added slide with image: {image_path.name}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to add image slide {image_path}: {e}")
            return False
    
    def _get_image_dimensions(self, image_path: Path) -> Tuple[int, int]:
        """
        Get image dimensions
        
        Args:
            image_path: Path to image file
            
        Returns:
            Tuple of (width, height) in pixels
        """
        try:
            with Image.open(image_path) as img:
                return img.size
        except Exception as e:
            logger.warning(f"Failed to get image dimensions for {image_path}: {e}")
            return (1920, 1080)  # Default dimensions
    
    def _calculate_image_position(self, img_width: int, img_height: int) -> Tuple[float, float, float, float]:
        """
        Calculate optimal image position and size on slide
        
        Args:
            img_width: Image width in pixels
            img_height: Image height in pixels
            
        Returns:
            Tuple of (left, top, width, height) in inches
        """
        # Available space on slide
        available_width = PPTSettings.IMAGE_WIDTH
        available_height = PPTSettings.IMAGE_HEIGHT
        
        # Calculate aspect ratios
        img_aspect = img_width / img_height
        available_aspect = available_width / available_height
        
        # Calculate dimensions maintaining aspect ratio
        if img_aspect > available_aspect:
            # Image is wider - fit to width
            width = available_width
            height = width / img_aspect
        else:
            # Image is taller - fit to height
            height = available_height
            width = height * img_aspect
        
        # Center the image
        left = PPTSettings.IMAGE_LEFT + (available_width - width) / 2
        top = PPTSettings.IMAGE_TOP + (available_height - height) / 2
        
        return left, top, width, height
    
    def _add_slide_title(self, slide, title: str):
        """
        Add title to slide
        
        Args:
            slide: PowerPoint slide object
            title: Title text
        """
        try:
            # Add title text box
            left = Inches(0.5)
            top = Inches(0.1)
            width = Inches(PPTSettings.SLIDE_WIDTH - 1)
            height = Inches(0.8)
            
            title_box = slide.shapes.add_textbox(left, top, width, height)
            title_frame = title_box.text_frame
            title_frame.text = title
            
            # Format title
            paragraph = title_frame.paragraphs[0]
            paragraph.font.size = Pt(24)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(0, 0, 0)  # Black
            
        except Exception as e:
            logger.warning(f"Failed to add slide title: {e}")
    
    def create_presentation_from_images(self, image_paths: List[Path], 
                                      output_path: Path,
                                      add_titles: bool = False,
                                      title_prefix: str = "Slide") -> bool:
        """
        Create complete presentation from list of images
        
        Args:
            image_paths: List of image file paths
            output_path: Path to save the presentation
            add_titles: Whether to add titles to slides
            title_prefix: Prefix for slide titles
            
        Returns:
            True if successful
        """
        if not image_paths:
            logger.error("No images provided for presentation")
            return False
        
        logger.info(f"Creating presentation with {len(image_paths)} slides")
        
        # Remove default slide if it exists and is empty
        if len(self.presentation.slides) == 1:
            slide = self.presentation.slides[0]
            if len(slide.shapes) == 0:
                # Remove empty slide
                slide_id = slide.slide_id
                self.presentation.part.drop_rel(slide.part.partname)
                del self.presentation.slides._sldIdLst[0]
        
        # Add slides for each image
        progress = ProgressTracker(len(image_paths), "Creating slides")
        successful_slides = 0
        
        for i, image_path in enumerate(image_paths):
            # Generate slide title if requested
            slide_title = None
            if add_titles:
                slide_title = f"{title_prefix} {i + 1}"
            
            # Add slide
            if self.add_image_slide(image_path, slide_title):
                successful_slides += 1
            
            progress.update()
        
        if successful_slides == 0:
            logger.error("No slides were successfully created")
            return False
        
        # Save presentation
        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self.presentation.save(str(output_path))
            
            # Get file size
            file_size = output_path.stat().st_size
            
            logger.info(f"Presentation saved: {output_path.name}")
            logger.info(f"File size: {format_file_size(file_size)}")
            logger.info(f"Total slides: {successful_slides}")
            
            return True
            
        except Exception as e:
            logger.error(f"Failed to save presentation: {e}")
            return False
    
    def add_title_slide(self, title: str, subtitle: Optional[str] = None) -> bool:
        """
        Add a title slide to the presentation
        
        Args:
            title: Main title text
            subtitle: Subtitle text (optional)
            
        Returns:
            True if successful
        """
        try:
            # Use title slide layout
            title_slide_layout = self.presentation.slide_layouts[0]
            slide = self.presentation.slides.add_slide(title_slide_layout)
            
            # Set title
            title_shape = slide.shapes.title
            title_shape.text = title
            
            # Set subtitle if provided
            if subtitle and len(slide.placeholders) > 1:
                subtitle_shape = slide.placeholders[1]
                subtitle_shape.text = subtitle
            
            logger.info(f"Added title slide: {title}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to add title slide: {e}")
            return False
    
    def add_section_slide(self, section_title: str) -> bool:
        """
        Add a section divider slide
        
        Args:
            section_title: Section title text
            
        Returns:
            True if successful
        """
        try:
            # Use section header layout or title layout
            layout_index = 2 if len(self.presentation.slide_layouts) > 2 else 0
            section_layout = self.presentation.slide_layouts[layout_index]
            slide = self.presentation.slides.add_slide(section_layout)
            
            # Set title
            title_shape = slide.shapes.title
            title_shape.text = section_title
            
            # Style the slide differently (e.g., different background)
            # This is optional and depends on the template
            
            logger.info(f"Added section slide: {section_title}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to add section slide: {e}")
            return False
    
    def get_slide_count(self) -> int:
        """
        Get current number of slides in presentation
        
        Returns:
            Number of slides
        """
        return len(self.presentation.slides)
    
    def set_presentation_properties(self, title: Optional[str] = None,
                                  author: Optional[str] = None,
                                  subject: Optional[str] = None,
                                  comments: Optional[str] = None):
        """
        Set presentation metadata properties
        
        Args:
            title: Presentation title
            author: Author name
            subject: Subject/topic
            comments: Comments/description
        """
        try:
            core_props = self.presentation.core_properties
            
            if title:
                core_props.title = title
            if author:
                core_props.author = author
            if subject:
                core_props.subject = subject
            if comments:
                core_props.comments = comments
            
            # Set creation date
            core_props.created = time.time()
            core_props.modified = time.time()
            
            logger.info("Set presentation properties")
            
        except Exception as e:
            logger.warning(f"Failed to set presentation properties: {e}")
    
    def save_presentation(self, output_path: Path) -> bool:
        """
        Save the presentation to file
        
        Args:
            output_path: Path to save the presentation
            
        Returns:
            True if successful
        """
        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self.presentation.save(str(output_path))
            
            file_size = output_path.stat().st_size
            logger.info(f"Presentation saved: {output_path.name}")
            logger.info(f"File size: {format_file_size(file_size)}")
            
            return True
            
        except Exception as e:
            logger.error(f"Failed to save presentation: {e}")
            return False
